"""ECS Pro - แปลง Export Summary Report (xlsx) ของ TONG HEER FASTENERS (THAILAND) เป็น INVB_*.xlsx สำหรับนำเข้า ECS

ไฟล์นี้คือส่วนคำนวณ (อ่านรายงาน / ใส่กฎ / เขียนไฟล์) ส่วนหน้าจออยู่ใน tongheer_app.pyw
ใช้แบบคำสั่งได้:  python tongheer_invb.py "รายงาน.xlsx" [ไฟล์ผลลัพธ์.xlsx] [--no-price]   (ค่าเริ่มต้นใส่คอลัมน์ PRICE)
"""
import datetime as dt
import json
import os
import re
import shutil
import sys
import tempfile
import urllib.request

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

VERSION = "1.1.3"  # เปลี่ยนเลขนี้พร้อมกับ version.json ทุกครั้งที่ออกเวอร์ชันใหม่
# เวอร์ชันล่าสุดอ่านจาก branch main บน GitHub (เครื่องที่ติดตั้งแล้วจะถามอัปเดตเองตอนเปิดโปรแกรม)
UPDATE_URL = "https://raw.githubusercontent.com/TPW0305/Draftongheer/main/"

# โฟลเดอร์ของโปรแกรม: ถ้าเป็น .exe ใช้โฟลเดอร์ที่ .exe อยู่ (json / DATA วางข้างกัน)
APP_DIR = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
CFG_PATH = os.path.join(APP_DIR, "tongheer_map.json")
DATA_PATH = os.path.join(APP_DIR, "DATA_TONGHEER.xlsx")  # product master จาก ECS (ไม่มีก็ใช้งานได้)

HEADERS = (
    "EXP_TAXNO,EXP_BRANCH,INV_NO,INV_DATE,INCOTERM,PAYTERM,CONSIGNEE,PONO,ITEM,PDTGROUP,PARTNO,PARTNAME,PARTNAME2,"
    "PARTNAME3,PARTNAMET,PARTNAMET2,PARTNAMET3,BRAND,MARK1,MARK2,MARK3,MARK4,TARIFF,SE,STATISTIC,PRIVILEGE,"
    "IMPORTTARRIF,QUANTITY,QUNIT,TARIFFQTY,TARIFFUNITY,GROSS,WEIGHT,WEIGHTU,INVCURR,AMOUNT,NUMPKG,UNTPKG,PRODUCTYEAR,"
    "FOC,FORMULAR1,F19,RAWCODE,BOI,BOILIC,FREEZONE,EPZ,IMPTAXINC,ORIGIN,DREMARK,MODELNO,MODELVERSION,MODELTAXNO,"
    "RTCPRODUCTCODE,PERMIT1,REFNO,ITEMSEP,SUBGROUP"
).split(",")
# คอลัมน์ที่ต้องเป็น Text (รักษาเลขศูนย์นำหน้า)
TEXT_COLS = {"EXP_TAXNO", "EXP_BRANCH", "INV_NO", "CONSIGNEE", "PDTGROUP", "PARTNO", "TARIFF", "SE", "STATISTIC",
             "PRIVILEGE", "RAWCODE", "BOI", "PERMIT1", "SUBGROUP"}
MONEY_COLS = {"QUANTITY", "TARIFFQTY", "GROSS", "WEIGHT", "AMOUNT"}
# ความกว้างคอลัมน์ตามไฟล์ INVB ตัวอย่าง (ที่เหลือ = 12)
WIDE = {3: 26, 12: 52, 13: 30, 15: 30, 19: 40, 27: 14, 31: 13, 39: 13, 52: 14, 54: 16}

# ค่าเริ่มต้น: เขียนลง tongheer_map.json ครั้งแรกที่รัน หลังจากนั้นแก้ที่ไฟล์ json ได้เลย
_PN = lambda prefix, partno, subgroup=None, ok=True: {  # noqa: E731
    "prefix": prefix, "partno": partno, "subgroup": subgroup or partno, "subConfirmed": ok}
_PD = lambda name, model, partno: {"prefix": "2THA", "name": name, "model": model, "partno": partno}  # noqa: E731
DEFAULT_CFG = {
    "constants": {
        "taxno": "0205547018706", "incoterm": "FOB", "brand": '"NO BRAND"', "tariff": "000073181590", "se": "50001",
        "statistic": "000", "privilege": "003", "importtarrif": "9PART3", "unit": "KGM", "currency": "USD",
        "pkgunit": "PX", "productyear": 2026, "origin": "TH", "partnamet": "สกรูทำด้วยสเตนเลส", "pdtgroup": "02",
        "permit1": "BOI0205547018706|0994000740964|05/10/2026", "itemsep": "N",
    },
    "prefixes": [
        {"prefix": "1THA", "branch": "0000", "boi": "1366(2)/2548"},
        {"prefix": "2THA", "branch": "0001", "boi": "66-1786-2-04-1-0"},
        {"prefix": "3THA", "branch": "0001", "boi": "61-1365-1-00-0-0"},
    ],
    "consignees": [
        {"mark": "WURTH INDUSTRY USA (REVCAR FASTENERS)", "code": "WURTH RE"},
        {"mark": "BBI", "code": "TONG WIN"},
        {"mark": "FASTCO", "code": "FASTENAL"},
        {"mark": "S.S.C", "code": "STAR"},
    ],
    # ลูกค้าที่ใช้ MARK1 ตามชื่อลูกค้า (ไม่ใช้ Marking ท้ายรายงาน)
    "customer_marks": [
        {"customer": "LAWSON PRODUCTS,INC", "mark": "LAWSON"},
        {"customer": "HARDWARE SPECIALTY CO.,INC.", "mark": "HARDWARE"},
    ],
    "partnos": [
        _PN("1THA", "304 HEX (AM)", ok=False), _PN("1THA", "316 HE(AM)", ok=False),
        _PN("1THA", "304 HEX LAG A2  (AM)", "304 HEX LAG A2"), _PN("1THA", "304 DIN 931(AM)", ok=False),
        _PN("1THA", "316 DIN 931(AM)", "316 DIN 931"), _PN("1THA", "304 DIN 933 (AM)", ok=False),
        _PN("1THA", "DIN 912 SOCKET 316(AM)", ok=False), _PN("1THA", "DIN 912 SOCKET304 (AM)", ok=False),
        _PN("1THA", "933 DIN316 (AM)", ok=False),
        _PN("2THA", "304 SOCKET(PT)"), _PN("2THA", "316 SOCKET(PT)"), _PN("2THA", "304 BUTTON SOCKET(PT)"),
        _PN("2THA", "316 BUTTON SOCKET(PT)"), _PN("2THA", "304  SOCKET FLAT (PT)", "304 SOCKET FLAT(PT)"),
        _PN("2THA", "316  SOCKET FLAT (PT)", "316 SOCKET FLAT(PT)"), _PN("2THA", "304 HEX CAR  (PT)", "304 HEX CAR(PT)"),
        _PN("2THA", "316 HEX CAR  (PT)", "316 HEX CAR(PT)"), _PN("2THA", "304 HEX (PT)"), _PN("2THA", "316 HE(PT)"),
        _PN("2THA", "304 DIN 933 (PT)"), _PN("2THA", "304 HEX FLANGE (PT)"), _PN("2THA", "316 (BOLT PT)"),
        _PN("2THA", "304 THREA(PT)"),
        _PN("3THA", "304 HEX (PT3)", ok=False),
    ],
    "products": [
        _PD("HEX. SOCKET SCREWS STAINLESS STEEL 304", "HEX SOCKET 304", "304 SOCKET(PT)"),
        _PD("HEX. SOCKET SCREWS STAINLESS STEEL 316", "HEX SOCKET 316", "316 SOCKET(PT)"),
        _PD("HEX. BUTTON SOCKET SCREWS STAINLESS STEEL 304", "HEX. BUTTON 304", "304 BUTTON SOCKET(PT)"),
        _PD("HEX. CARRIAGE BOLTS STAINLESS STEEL 316", "HEX CARRIAGE 316", "316 HEX CAR  (PT)"),
        _PD("HEX BOLTS ( NO WASHER FACE ) STAINLESS STEEL 316", "HEX BOLTS 316", "316 (BOLT PT)"),
        _PD("HEX CAP SCREWS STAINLESS STEEL 304", "HEX CAP 304", "304 HEX (PT)"),
        _PD("HEX CAP SCREWS STAINLESS STEEL 316", "HEX CAP 316", "316 HE(PT)"),
    ],
}


# ---------------- ตัวช่วย ----------------
def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


def key_of(s):
    return re.sub(r"[^A-Za-z0-9]", "", str(s or "")).upper()


def fmt2(v):
    return "-" if v is None else f"{v:,.2f}"


def date_text(d):
    return d.strftime("%d/%m/%Y") if d else ""


def parse_date_text(text):
    """รับ dd/mm/yyyy (ปี ค.ศ.) คืน date หรือ None"""
    m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*", text or "")
    if not m:
        return None
    try:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


# ---------------- config ----------------
def load_config():
    if not os.path.exists(CFG_PATH):
        save_config(DEFAULT_CFG)
    with open(CFG_PATH, encoding="utf-8-sig") as f:
        cfg = json.load(f)
    if "customer_marks" not in cfg:  # json จากเวอร์ชันก่อน ECS Pro ยังไม่มีกฎ mark ตามลูกค้า
        cfg["customer_marks"] = [dict(c) for c in DEFAULT_CFG["customer_marks"]]
        save_config(cfg)
    return cfg


def save_config(cfg):
    with open(CFG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def find_consignee(cfg, mark):
    k = norm(mark)
    return next((c["code"] for c in cfg["consignees"] if norm(c["mark"]) == k), "")


def find_customer_mark(cfg, customer):
    """MARK1 ที่กำหนดตามชื่อลูกค้า (เทียบเฉพาะตัวอักษร/ตัวเลข ไม่สนจุด จุลภาค ช่องว่าง) หรือ "" ถ้าไม่มีกฎ"""
    k = key_of(customer)
    return next((c["mark"] for c in cfg.get("customer_marks", []) if k and key_of(c["customer"]) == k), "")


def find_product(cfg, prefix, name, model):
    n, m = norm(name), norm(model)
    return next((p for p in cfg["products"]
                 if p["prefix"] == prefix and norm(p["name"]) == n and norm(p["model"]) == m), None)


def find_partno(cfg, prefix, partno):
    return next((p for p in cfg["partnos"] if p["prefix"] == prefix and p["partno"] == partno), None)


def default_subgroup(cfg, prefix, partno):
    p = find_partno(cfg, prefix, partno)
    return p["subgroup"] if p else partno


def known_partnos(cfg, prefix):
    return [p["partno"] for p in cfg["partnos"] if p["prefix"] == prefix]


# ---------------- product master (DATA_TONGHEER.xlsx) ----------------
def load_data(path=DATA_PATH):
    """คืน [(Code, DecDesc1 แบบ norm, BOINo)] จาก Sheet1 หรือ [] ถ้าไม่มีไฟล์/อ่านไม่ได้"""
    if not os.path.exists(path):
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
        h = {str(n): i for i, n in enumerate(rows[0]) if n}
        out = []
        for r in rows[1:]:
            code, desc, boi = r[h["Code"]], r[h["DecDesc1"]], r[h["BOINo"]]
            if code and desc:
                out.append((str(code), norm(str(desc).replace("_x000D_", " ")), str(boi or "").strip()))
        return out
    except Exception:
        return []


def data_candidates(data, boi, name, model_line):
    """รหัสใน DATA ที่ DecDesc1 (ชื่อ + บรรทัด MODEL) และ BOI ตรงกับรายการ"""
    want = norm(f"{name} {model_line}")
    return [code for code, desc, b in data if b == boi and desc == want]


# ---------------- อ่านรายงาน ----------------
class Item:
    def __init__(self, name, model_line, rawcode, net, gross, amount):
        self.name, self.model_line, self.rawcode = name, model_line, rawcode
        self.net, self.gross, self.amount = net, gross, amount
        self.partno = ""
        self.subgroup = ""
        self.candidates = []  # รหัสที่ตรงใน DATA_TONGHEER (ใช้เมื่อยังไม่เคยยืนยัน)


class Report:
    def __init__(self):
        self.path = ""
        self.invoices = []
        self.date = None
        self.etd = None
        self.items = []
        self.tot_net = self.tot_gross = self.tot_amount = None
        self.pallets = None
        self.marking = ""
        self.customer = ""
        self.warnings = []
        # เติมโดย resolve()
        self.inv_no = self.prefix = self.branch = self.boi = self.consignee = ""

    def sums(self):
        return (round(sum(i.net for i in self.items), 2), round(sum(i.gross for i in self.items), 2),
                round(sum(i.amount for i in self.items), 2))

    def total_checks(self):
        """คืน [(ชื่อ, ผลรวมรายการ, ยอดในรายงาน, ส่วนต่าง หรือ None)]"""
        out = []
        for label, s, ref in zip(("Net", "Gross", "Amount"), self.sums(), (self.tot_net, self.tot_gross, self.tot_amount)):
            out.append((label, s, ref, None if ref is None else round(s - ref, 2)))
        return out

    def totals_text(self):
        parts = []
        for label, s, ref, diff in self.total_checks():
            if ref is None:
                parts.append(f"{label} {fmt2(s)} (ไม่พบยอดในรายงาน)")
            elif abs(diff) < 0.005:
                parts.append(f"{label} {fmt2(s)} = รายงาน")
            else:
                parts.append(f"{label} {fmt2(s)} / รายงาน {fmt2(ref)} (ต่าง {diff:+.2f})")
        return f"ยอดรวม {len(self.items)} รายการ: " + " | ".join(parts)


def _as_date(v, month_first=True):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.match(r"\s*(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})", str(v or ""))
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y += 2000 if y < 100 else 0
        for x, z in ((a, b), (b, a)) if month_first else ((b, a), (a, b)):  # ถ้าอ่านแบบเดือนก่อนไม่ได้ (เช่น 17-10-2026) ลองวันก่อน
            try:
                return dt.date(y, x, z)
            except ValueError:
                continue
    return None


def parse_report(path):
    if os.path.splitext(path)[1].lower() != ".xlsx":
        raise ValueError("รองรับเฉพาะไฟล์ .xlsx (ถ้าเป็น .xls ให้ Save As เป็น .xlsx ก่อน)")
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    cells = {}  # (row, col) -> value
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                cells[(c.row, c.column)] = c.value
    wb.close()

    rep = Report()
    rep.path = path
    warn = rep.warnings
    texts = [(rc, v) for rc, v in sorted(cells.items()) if isinstance(v, str)]

    def find(pred):
        return next((rc for rc, v in texts if pred(key_of(v))), None)

    def right(rc):
        if not rc:
            return None
        cands = [c for (r, c) in cells if r == rc[0] and c > rc[1]]
        return (rc[0], min(cands)) if cands else None

    def label_value(label):
        """ค่าหลังหัวข้อ เช่น "Marking :" อยู่เซลล์ถัดไป หรือบางรายงานอยู่เซลล์เดียวกัน ("Marking :WURTH ...")"""
        rc = find(lambda k: k == label)
        if rc:
            return str(cells.get(right(rc), "")).strip()
        rc = find(lambda k: k.startswith(label))
        if not rc:
            return ""
        inline = re.sub(r"^\s*" + label + r"\s*:?\s*", "", cells[rc], flags=re.I).strip()
        return inline or str(cells.get(right(rc), "")).strip()

    def num(rc):
        v = cells.get(rc) if rc else None
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    if not find(lambda k: k.startswith("EXPORTSUMMARYREPORT")):
        warn.append('ไม่พบหัวรายงาน "Export Summary Report" - ตรวจว่าเลือกไฟล์ถูกหรือไม่')

    # --- เลขอินวอย: ใช้ตารางสรุปตามอินวอยด้านล่างถ้ามี ไม่งั้นใช้ช่อง Invoice# ---
    rx_inv = re.compile(r"\s*([1-3]THA\d+)\s*$")
    inv_hdr = find(lambda k: k == "INVOICENO")
    if inv_hdr:
        r = inv_hdr[0] + 1
        while True:
            m = rx_inv.match(str(cells.get((r, inv_hdr[1]), "")))
            if not m:
                break
            rep.invoices.append(m.group(1))
            r += 1
        amt_col = next((c for (rr, c), v in texts if rr == inv_hdr[0] and key_of(v) == "AMOUNT"), None)
        tot_row = next((rr for (rr, c), v in texts if c == inv_hdr[1] and rr > inv_hdr[0] and key_of(v) == "TOTAL"), None)
        if amt_col and tot_row:
            rep.tot_amount = num((tot_row, amt_col))
    inv_lbl = find(lambda k: k.startswith("INVOICE") and k != "INVOICENO")
    inv_cell = right(inv_lbl)
    if not rep.invoices:
        m = rx_inv.match(str(cells.get(inv_cell, ""))) if inv_cell else None
        if m:
            rep.invoices.append(m.group(1))
        else:
            warn.append("ไม่พบเลขอินวอยในรายงาน - กรอก INV_NO เอง")

    # --- วันที่อินวอย / ETD ---
    if inv_lbl:
        for (r, c) in sorted(cells):
            if r == inv_lbl[0] and c > inv_lbl[1] and (r, c) != inv_cell:
                d = _as_date(cells[(r, c)])
                if d:
                    rep.date = d
                    break
    etd_cell = right(find(lambda k: k == "ETD"))
    eta_cell = right(find(lambda k: k == "ETA"))
    if not rep.date:  # รายงานแบบใหม่: Invoice# อยู่ท้ายรายงาน วันที่อยู่หัวรายงาน (ก่อนตารางสินค้า) ไม่ใช่แถวเดียวกับ Invoice#
        net_row = (find(lambda k: k == "NETWEIGHT") or (10**6,))[0]
        for (r, c) in sorted(cells):
            if r < net_row and (r, c) not in (etd_cell, eta_cell):
                d = _as_date(cells[(r, c)])
                v = cells[(r, c)]  # ข้อความรับเฉพาะที่เป็นวันที่ทั้งช่อง เช่น "15-10-2026" (ไม่หยิบที่อยู่/เลขเอกสาร)
                if d and (not isinstance(v, str) or re.fullmatch(r"\s*\d{1,2}[-/]\d{1,2}[-/]\d{2,4}(\s+[\d:]+)?\s*", v)):
                    rep.date = d
                    break
    rep.etd = _as_date(cells.get(etd_cell)) if etd_cell else None
    if not rep.date:
        warn.append("ไม่พบวันที่อินวอยในรายงาน - กรอก INV_DATE เอง (dd/mm/yyyy)")
    elif not rep.etd:
        warn.append("ไม่พบ ETD จึงตรวจการสลับวัน/เดือนของ INV_DATE ไม่ได้")
    else:
        gap = (rep.etd - rep.date).days
        if gap < 0 or gap > 60:
            msg = f"INV_DATE {date_text(rep.date)} ไม่สอดคล้องกับ ETD {date_text(rep.etd)}"
            d = rep.date
            if d.day <= 12 and d.day != d.month:
                sw = dt.date(d.year, d.day, d.month)
                if 0 <= (rep.etd - sw).days <= 60:  # สลับแล้วสอดคล้องกับ ETD: แก้ให้เอง แล้วแจ้งผู้ใช้
                    rep.date = sw
                    warn.append(f"INV_DATE ในรายงานเป็น {date_text(d)} สลับวัน/เดือน - แก้ให้เป็น {date_text(sw)} "
                                f"ตาม ETD {date_text(rep.etd)} (ตรวจอีกครั้งในช่อง INV_DATE)")
                    msg = None
            if msg:
                warn.append(msg + " (ตรวจและแก้ในช่อง INV_DATE)")

    # --- รายการสินค้า (ตารางส่วนบน ไม่ใช่ตารางสรุปตามอินวอย) ---
    net_h = find(lambda k: k == "NETWEIGHT")
    if not net_h:
        raise ValueError("ไม่พบหัวคอลัมน์ NetWeight ในรายงาน")
    hdr = {key_of(v): c for (r, c), v in texts if r == net_h[0]}
    if "GROSSWEIGHT" not in hdr or "AMOUNT" not in hdr:
        raise ValueError("ไม่พบหัวคอลัมน์ GrossWeight / Amount ในรายงาน")
    desc_col = hdr.get("SUBGROUP", 3)
    sub_rows = sorted({r for (r, c), v in texts if key_of(v).startswith("SUBTOTALFOR")})
    if not sub_rows:
        raise ValueError('ไม่พบรายการสินค้า (บรรทัด "Sub Total For :") ในรายงาน')
    prev = net_h[0]
    for r in sub_rows:
        lines = [re.sub(r"\s*[\r\n]+\s*", " ", str(cells[(rr, desc_col)])).strip()
                 for rr in range(prev + 1, r) if (rr, desc_col) in cells]
        name = model_line = raw = ""
        for line in lines:
            m = re.match(r"\(\s*MODEL\s*:\s*(.*?)\s*\)\s*$", line)
            if m:
                if not model_line:
                    model_line, raw = line, m.group(1)
            elif not name:
                name = line
        n = len(rep.items) + 1
        net, gross, amount = num((r, net_h[1])), num((r, hdr["GROSSWEIGHT"])), num((r, hdr["AMOUNT"]))
        if None in (net, gross, amount):
            raise ValueError(f"รายการ {n} (แถว {r}): อ่าน Net/Gross/Amount ไม่ได้")
        if not name:
            warn.append(f"รายการ {n}: ไม่พบชื่อสินค้า")
        if not model_line:
            warn.append(f"รายการ {n}: ไม่พบบรรทัด ( MODEL : ... ) - RAWCODE จะว่าง")
        rep.items.append(Item(name, model_line, raw, round(net, 2), round(gross, 2), round(amount, 2)))
        prev = r

    # --- ยอดรวม / พาเลท / marking ---
    rep.tot_net = num(right(find(lambda k: k.startswith("TOTALNETWEIGHT"))))
    rep.tot_gross = num(right(find(lambda k: k.startswith("TOTALGROSSWEIGHT"))))
    if rep.tot_amount is None:  # ไม่มีตารางสรุปตามอินวอย: ใช้ช่อง "Total Amount :" (หรือ FinalAmount) ท้ายรายงาน
        rep.tot_amount = num(right(find(lambda k: k.startswith("TOTALAMOUNT")))) or num(right(find(lambda k: k.startswith("FINALAMOUNT"))))
    if rep.tot_amount is None:  # ไม่มีทั้งสองช่อง: ใช้แถวผลรวมใต้รายการสุดท้าย
        for r in range(prev + 1, prev + 4):
            v = num((r, hdr["AMOUNT"]))
            if v is not None:
                rep.tot_amount = v
                break
    p = num(right(find(lambda k: k.startswith("TOTALPALLET"))))
    rep.pallets = int(p) if p is not None else None
    if rep.pallets is None:
        warn.append("ไม่พบ Total Pallet - กรอก NUMPKG เอง")
    rep.marking = label_value("MARKING")
    rep.customer = label_value("CUSTOMER")
    return rep


# ---------------- ใส่กฎของ skill ----------------
def resolve(rep, cfg, data=None):
    """เติม inv_no / branch / boi / consignee / partno / subgroup ตามกฎ และเพิ่มคำเตือน"""
    warn = rep.warnings
    if rep.invoices:
        first = rep.invoices[0]
        rep.prefix = first[:4]
        rep.inv_no = first
        if len(rep.invoices) == 2:
            rep.inv_no = ",".join(rep.invoices)
            warn.append(f"รายงานมี 2 อินวอย - ใส่ INV_NO เลขเต็มทั้งสองใบ {rep.inv_no} และรวมเป็นไฟล์เดียว")
        elif len(rep.invoices) > 2:
            rep.inv_no = first + "," + ",".join(i[-3:] for i in rep.invoices[1:])
            warn.append(f"รายงานมี {len(rep.invoices)} อินวอย - ใส่ INV_NO แบบย่อ {rep.inv_no}")
        if any(i[:4] != rep.prefix for i in rep.invoices):
            warn.append(f"อินวอยในรายงานมี prefix ไม่เหมือนกัน - ใช้กฎของ {rep.prefix} ทั้งไฟล์")
        pf = next((p for p in cfg["prefixes"] if p["prefix"] == rep.prefix), None)
        if pf:
            rep.branch, rep.boi = pf["branch"], pf["boi"]
        else:
            warn.append(f"ไม่รู้จัก prefix {rep.prefix} (ไม่มี EXP_BRANCH / BOI)")
    cust_mark = find_customer_mark(cfg, rep.customer)
    if cust_mark:
        if rep.marking and norm(rep.marking) != norm(cust_mark):
            warn.append(f'ลูกค้า "{rep.customer}" ใช้ MARK1 "{cust_mark}" ตามกฎลูกค้า (แทน Marking ในรายงาน "{rep.marking}")')
        rep.marking = cust_mark
    elif not rep.marking:
        warn.append("ไม่พบ Marking ท้ายรายงาน - กรอก MARK1 เอง")
    rep.consignee = find_consignee(cfg, rep.marking)
    if rep.marking and not rep.consignee:
        warn.append(f'Marking "{rep.marking}" ยังไม่เคยกำหนด CONSIGNEE - กรอกรหัสตาม ECS (โปรแกรมจะจำไว้)')
    data = load_data() if data is None else data
    suffix = {"1THA": "(AM)", "2THA": "(PT)", "3THA": "(PT3)"}.get(rep.prefix, "")
    for n, it in enumerate(rep.items, 1):
        p = find_product(cfg, rep.prefix, it.name, it.rawcode)
        if p:  # เคยยืนยันแล้ว
            it.partno = p["partno"]
        else:  # ยังไม่เคยยืนยัน: เสนอจาก DATA_TONGHEER ตามชื่อ + BOI
            it.candidates = data_candidates(data, rep.boi, it.name, it.model_line) if rep.boi else []
            if it.candidates:
                confirmed = [c for c in it.candidates if find_partno(cfg, rep.prefix, c)]
                by_suffix = [c for c in it.candidates if suffix and c.rstrip().endswith(suffix)]
                it.partno = (confirmed or by_suffix or it.candidates)[0]
                msg = f'รายการ {n}: PARTNO "{it.partno}" เลือกจาก DATA_TONGHEER ยังไม่เคยยืนยัน - ตรวจก่อนใช้'
                if len(it.candidates) > 1:
                    msg += " (ใน DATA มีชื่อซ้ำ: " + ", ".join(f'"{c}"' for c in it.candidates) + ")"
                warn.append(msg)
        if it.partno:
            it.subgroup = default_subgroup(cfg, rep.prefix, it.partno)
    return rep


def subgroup_warnings(rep, cfg):
    out = []
    for n, it in enumerate(rep.items, 1):
        if it.partno:
            p = find_partno(cfg, rep.prefix, it.partno)
            if not p or not p.get("subConfirmed"):
                out.append(f'รายการ {n}: SUBGROUP "{it.subgroup}" ยังไม่เคยยืนยันกับหน้าจอ ECS - เช็กก่อนใช้')
    return out


# ---------------- เขียน INVB ----------------
def permit1_text(base, issue_date=None):
    """PERMIT1 = เลขที่|ผู้ออก|วันที่ออก  (ส่วนที่ 3 เปลี่ยนตามที่เลือก ถ้าไม่เลือกใช้วันที่ปัจจุบัน)"""
    parts = str(base).split("|")
    parts = (parts + ["", ""])[:2] if len(parts) < 3 else parts[:2]
    return "|".join(parts + [date_text(issue_date or dt.date.today())])


def write_invb(rep, cfg, out_path, inv_no, inv_date, mark, consignee, pallets, with_price=False, permit_date=None):
    k = cfg["constants"]
    permit1 = permit1_text(k["permit1"], permit_date)
    headers = HEADERS + (["PRICE"] if with_price else [])
    wb = Workbook()
    ws = wb.active
    ws.title = "SHEET1"
    f_head, f_body = Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10)
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h).font = f_head
    for n, it in enumerate(rep.items, 1):
        v = {
            "EXP_TAXNO": k["taxno"], "EXP_BRANCH": rep.branch, "INV_NO": inv_no, "INV_DATE": inv_date,
            "INCOTERM": k["incoterm"], "CONSIGNEE": consignee, "ITEM": n, "PDTGROUP": k["pdtgroup"],
            "PARTNO": it.partno, "PARTNAME": it.name, "PARTNAME2": it.model_line, "PARTNAMET": k["partnamet"],
            "BRAND": k["brand"], "MARK1": mark if n == 1 else "-DO-", "TARIFF": k["tariff"], "SE": k["se"],
            "STATISTIC": k["statistic"], "PRIVILEGE": k["privilege"], "IMPORTTARRIF": k["importtarrif"],
            "QUANTITY": it.net, "QUNIT": k["unit"], "TARIFFQTY": it.net, "TARIFFUNITY": k["unit"], "GROSS": it.gross,
            "WEIGHT": it.net, "WEIGHTU": k["unit"], "INVCURR": k["currency"], "AMOUNT": it.amount,
            "NUMPKG": pallets if n == 1 else 0, "UNTPKG": k["pkgunit"], "PRODUCTYEAR": k["productyear"],
            "RAWCODE": it.rawcode, "BOI": rep.boi, "ORIGIN": k["origin"], "PERMIT1": permit1,
            "ITEMSEP": k["itemsep"], "SUBGROUP": it.subgroup,
        }
        if with_price:
            v["PRICE"] = it.amount / it.net
        for c, h in enumerate(headers, 1):
            val = v.get(h)
            if val is None or val == "":
                continue
            cell = ws.cell(n + 1, c, val)
            cell.font = f_body
            if h in TEXT_COLS:
                cell.value = str(val)
                cell.number_format = "@"
            elif h == "INV_DATE":
                cell.number_format = "dd/mm/yyyy"
            elif h in MONEY_COLS:
                cell.number_format = "#,##0.00"
    for c in range(1, 56):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = WIDE.get(c, 12)
    ws.freeze_panes = "A2"
    wb.save(out_path)


def learn(rep, cfg, mark, consignee):
    """จำ CONSIGNEE / PARTNO / SUBGROUP ที่ผู้ใช้กรอก ลง tongheer_map.json"""
    changed = False
    if mark and consignee:
        hit = next((c for c in cfg["consignees"] if norm(c["mark"]) == norm(mark)), None)
        if not hit:
            cfg["consignees"].append({"mark": mark, "code": consignee})
            changed = True
        elif hit["code"] != consignee:
            hit["code"] = consignee
            changed = True
    for it in rep.items:
        p = find_product(cfg, rep.prefix, it.name, it.rawcode)
        if not p:
            cfg["products"].append({"prefix": rep.prefix, "name": it.name, "model": it.rawcode, "partno": it.partno})
            changed = True
        elif p["partno"] != it.partno:
            p["partno"] = it.partno
            changed = True
        pn = find_partno(cfg, rep.prefix, it.partno)
        if not pn:
            cfg["partnos"].append({"prefix": rep.prefix, "partno": it.partno, "subgroup": it.subgroup, "subConfirmed": True})
            changed = True
        elif pn["subgroup"] != it.subgroup:
            pn["subgroup"], pn["subConfirmed"] = it.subgroup, True
            changed = True
    if changed:
        save_config(cfg)


# ---------------- อัปเดตโปรแกรม ----------------
def _ver(text):
    return tuple(int(x) for x in re.findall(r"\d+", str(text or ""))[:3])


def _fetch(name, timeout):
    with urllib.request.urlopen(UPDATE_URL + urllib.request.quote(name), timeout=timeout) as r:
        return r.read()


def check_update(timeout=6):
    """คืน dict ของ version.json บน GitHub ถ้ามีเวอร์ชันใหม่กว่าเครื่องนี้ ไม่งั้นคืน None (ไม่มีเน็ต/อ่านไม่ได้ = None)"""
    if os.path.exists(os.path.join(APP_DIR, ".git")):  # โฟลเดอร์พัฒนา ไม่อัปเดตทับ
        return None
    try:
        info = json.loads(_fetch("version.json", timeout).decode("utf-8-sig"))
        return info if _ver(info["version"]) > _ver(VERSION) else None
    except Exception:
        return None


def apply_update(info, timeout=30):
    """ดาวน์โหลดไฟล์โปรแกรมเวอร์ชันใหม่มาวางทับ (ไม่แตะ tongheer_map.json ที่จำไว้) ถ้าผิดพลาดจะ raise และไม่เปลี่ยนไฟล์เดิม"""
    names = [n for n in info["files"] if os.path.basename(n) == n and n != os.path.basename(CFG_PATH)]
    tmp = tempfile.mkdtemp(prefix="ecspro_update_")
    try:
        for n in names:  # ดาวน์โหลดให้ครบและตรวจก่อน แล้วค่อยวางทับ
            data = _fetch(n, timeout)
            if n.endswith((".py", ".pyw")):
                compile(data, n, "exec")
            with open(os.path.join(tmp, n), "wb") as f:
                f.write(data)
        core_name = os.path.basename(__file__)
        if core_name in names:
            m = re.search(r'^VERSION = "([^"]+)"', open(os.path.join(tmp, core_name), encoding="utf-8").read(), re.M)
            if not m or _ver(m.group(1)) != _ver(info["version"]):
                raise ValueError("ไฟล์ที่ดาวน์โหลดไม่ตรงกับเวอร์ชันที่ประกาศ - ลองใหม่อีกครั้งภายหลัง")
        for n in names:
            shutil.copyfile(os.path.join(tmp, n), os.path.join(APP_DIR, n))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def out_path_for(rep, inv_no):
    return os.path.join(os.path.dirname(rep.path), f"INVB_{inv_no.split(',')[0].strip()}.xlsx")


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    if not args:
        print(__doc__)
        return 1
    cfg = load_config()
    rep = resolve(parse_report(os.path.abspath(args[0])), cfg)
    problems = [f"รายการ {n}: ยังไม่รู้ PARTNO ของ {it.name} {it.model_line}" for n, it in enumerate(rep.items, 1) if not it.partno]
    if not rep.consignee or not rep.branch or not rep.date or rep.pallets is None:
        problems.append("ข้อมูลหัวอินวอยไม่ครบ")
    for w in rep.warnings + problems:
        print("!", w)
    if problems:
        print("ข้อมูลยังไม่ครบ - เปิดโปรแกรมแบบมีหน้าจอ (tongheer_app.pyw) เพื่อกรอกส่วนที่ขาด")
        return 2
    out = args[1] if len(args) > 1 else out_path_for(rep, rep.inv_no)
    write_invb(rep, cfg, out, rep.inv_no, rep.date, rep.marking, rep.consignee, rep.pallets, "--no-price" not in argv)
    print("สร้างไฟล์:", out)
    print(rep.totals_text())
    for w in subgroup_warnings(rep, cfg):
        print("!", w)
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
